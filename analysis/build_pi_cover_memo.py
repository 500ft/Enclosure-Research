#!/usr/bin/env python3
"""Build a concise cover memo for the full PI literature synthesis."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "PI_Cover_Memo_Literature_and_Next_Steps.docx"
BLUE = "1F4E79"
DARK = "17324D"
LIGHT = "EAF2F8"
GRAY = "667085"
WHITE = "FFFFFF"


def run_style(run, size=10.5, bold=False, color="000000", italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def margins(cell, value=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        tc_mar.append(element)
    tc_pr.append(tc_mar)


def para(doc, text="", bold_lead=None, after=5, color="000000"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if bold_lead:
        run_style(p.add_run(bold_lead), bold=True, color=color)
        run_style(p.add_run(text[len(bold_lead):]), color=color)
    else:
        run_style(p.add_run(text), color=color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    run_style(p.add_run(text), size=10.2)


def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    for name, size in [("Heading 1", 15), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run_style(title.add_run("RESEARCH PROGRESS AND LITERATURE REVIEW MEMO"), size=18, bold=True, color=DARK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run_style(
        subtitle.add_run("Low-Cost Outdoor Multi-Sensor Box: Current Direction and Decisions Needed"),
        size=12,
        italic=True,
        color=GRAY,
    )

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    metadata = [
        ("To", "Principal Investigator"),
        ("Date", "June 11, 2026"),
        ("Current output", "Literature synthesis covering 19 sources"),
        ("Working paper", "Calibration, Ambient Conditions, and Integration Effects on the Accuracy and Autonomy of a Low-Cost Outdoor Multi-Sensor Box"),
    ]
    for row, (label, value) in zip(table.rows, metadata):
        row.cells[0].width = Inches(1.35)
        row.cells[1].width = Inches(5.0)
        shade(row.cells[0], LIGHT)
        for cell in row.cells:
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run_style(row.cells[0].paragraphs[0].add_run(label), size=9.5, bold=True, color=DARK)
        run_style(row.cells[1].paragraphs[0].add_run(value), size=9.5)

    heading(doc, "Purpose", 1)
    para(
        doc,
        "The literature review has been expanded and reorganized to support a system-level evaluation of the lab's existing outdoor sensor box. The attached full synthesis summarizes all 19 papers, including reported sensors, materials, geometries, experimental setups, findings, and implications for the proposed study.",
    )

    heading(doc, "What Changed", 1)
    for item in [
        "The project moved from an enclosure-only study to evaluation of the complete deployed sensor box.",
        "The literature set now covers radiation-shield geometry, enclosure materials, calibration, drift, autonomy, reliability, and long-term field failure.",
        "Each paper now has a dedicated sensor and physical-box pros/cons analysis.",
        "Material and geometry options have been converted into comparable engineering criteria rather than an unstructured list.",
        "The paper draft now frames calibration and physical integration as connected sources of measurement error.",
    ]:
        bullet(doc, item)

    heading(doc, "Strongest Literature Conclusions", 1)
    conclusions = [
        ("Airflow controls accuracy.", "Passive shields show their largest warm bias during high-solar, low-wind conditions. Active aspiration can improve accuracy but reduces autonomy and adds a fan failure mode."),
        ("The box is part of the instrument.", "Internal heat, sensor placement, inlet geometry, reflected radiation, moisture, and sealing affect measurements and cannot be treated as neutral packaging."),
        ("White, reflective, and weatherable materials are favored.", "White/off-white ASA is the strongest supported printed outdoor material. PLA is useful for prototyping but has durability limitations."),
        ("Autonomy is more than battery life.", "Connectors, mounts, storage, communication, firmware, charging, fans, and corrosion frequently require intervention before the sensor reaches end of life."),
        ("Calibration must be tested over time.", "A correction developed at one site or time period may not transfer. Pre/post deployment comparison and time-separated validation are important."),
    ]
    for label, text in conclusions:
        para(doc, f"{label} {text}", bold_lead=f"{label} ", after=4)

    doc.add_page_break()
    heading(doc, "Recommended Experimental Framing", 1)
    para(
        doc,
        "The proposed experiment should begin with the current lab box unchanged, then compare a small number of justified alternatives. The objective is to identify whether the box is limited mainly by calibration, enclosure design, power, firmware/data handling, or maintenance.",
    )
    designs = [
        ("Current box", "Baseline accuracy, autonomy, and failure behavior."),
        ("Strong passive option", "White reflective stacked-plate or truncated-cone shield with separated sealed electronics."),
        ("Optional active benchmark", "Aspirated shield used to quantify the maximum practical accuracy improvement and its power cost."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Design", "Role in experiment"]):
        shade(cell, BLUE)
        margins(cell)
        run_style(cell.paragraphs[0].add_run(text), size=9.2, bold=True, color=WHITE)
    for design, role in designs:
        cells = table.add_row().cells
        for cell in cells:
            margins(cell)
        run_style(cells[0].paragraphs[0].add_run(design), size=9.2, bold=True, color=DARK)
        run_style(cells[1].paragraphs[0].add_run(role), size=9.2)

    heading(doc, "Decisions Requested", 1)
    for item in [
        "Confirm which sensor variables and exact hardware models are in scope.",
        "Identify the reference instruments available for baseline co-location.",
        "Define acceptable calibrated error, minimum autonomy, and data-completeness thresholds.",
        "Confirm whether the first study should compare passive shields only or include an active aspirated benchmark.",
        "Confirm the intended deployment duration and representative weather conditions.",
    ]:
        bullet(doc, item)

    heading(doc, "Immediate Next Steps", 1)
    for item in [
        "Complete the baseline hardware and enclosure inventory.",
        "Prepare the current box for reference co-location without changing its design.",
        "Begin logging sensor data, reference data, battery status, environmental conditions, missing samples, and maintenance events.",
        "Use baseline results to select the smallest defensible set of enclosure variants for testing.",
    ]:
        bullet(doc, item)

    para(
        doc,
        "Attached report: PI_Literature_Synthesis_Outdoor_Sensor_Box.pdf",
        bold_lead="Attached report: ",
        after=0,
        color=BLUE,
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_style(footer.add_run("Outdoor Multi-Sensor Box | Research Progress Memo"), size=8.5, color=GRAY)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
