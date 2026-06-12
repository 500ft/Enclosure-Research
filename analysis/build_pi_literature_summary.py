#!/usr/bin/env python3
"""Build the PI-ready literature synthesis DOCX used to produce the final PDF."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_DOCX = OUT_DIR / "PI_Literature_Synthesis_Outdoor_Sensor_Box.docx"

BLUE = "1F4E79"
DARK = "17324D"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"
BLACK = "000000"


profiles = [
    {
        "category": "Complete Sensor-Box and Field-System Studies",
        "citation": "Theisen et al. (2020)",
        "title": "More Science with Less: Evaluation of a 3D-Printed Weather Station",
        "doi": "10.5194/amt-13-4699-2020",
        "sensors": "MCP9808; HTU21D; BMP280; SI1145; Hall-effect wind, direction, and rain sensors.",
        "materials": "Off-white/gray ASA; PVC frame; electrical junction box; polyurethane, PTFE filter, conformal coating, and silicone.",
        "geometry": "Integrated 3D-PAWS station; naturally aspirated shield; crossbar-mounted wind/UV sensors; separate junction box for logger.",
        "setup": "Eight-month outdoor deployment beside an Oklahoma Mesonet reference station.",
        "shows": "Several low-cost measurements can be useful, but wind dependence, corrosion, moisture, connectors, printed fasteners, and moving parts strongly affect total system performance.",
        "conclusion": "The sensing element is only one part of field accuracy and lifetime. Integration and maintenance failures can dominate.",
        "importance": "Provides the closest precedent for evaluating the lab box as a complete deployed instrument and for recording both accuracy and failure events.",
        "consider": [
            "Log internal box temperature and wind speed so radiation-shield performance can be separated from sensor calibration.",
            "Use metal structural fasteners and explicitly inspect connectors, drain paths, coatings, and exposed printed parts.",
        ],
    },
    {
        "category": "Complete Sensor-Box and Field-System Studies",
        "citation": "Tatsumi et al. (2021)",
        "title": "An Open-Source, Low-Cost Measurement System for Collecting Hydrometeorological Data in the Open Field",
        "doi": "10.3390/technologies9040078",
        "sensors": "LM60BIZ air/soil temperature; YL-69 soil moisture; G2711-01 GaAsP photodiode; GPS and micro-SD.",
        "materials": "3D-printed protective case using unspecified resin; handmade photodiode cover.",
        "geometry": "Compact rectangular logger case with external sensors connected through detachable terminals; solar-powered system.",
        "setup": "Grass-field deployment with reference comparisons; 30-day temperature and 19-day soil-moisture/PPFD measurements.",
        "shows": "An open-field system must be evaluated as a package of sensing, power, storage, enclosure, and external connections.",
        "conclusion": "Low-cost, modular hardware is practical, but exact material reporting and external-connection protection are necessary for reproducibility.",
        "importance": "Supports the baseline inventory and reinforces documenting the complete lab box before modification.",
        "consider": [
            "Record exact material chemistry rather than only stating that a case is 3D printed.",
            "Include connector sealing, strain relief, and external-probe replacement in maintainability testing.",
        ],
    },
    {
        "category": "Complete Sensor-Box and Field-System Studies",
        "citation": "Grimsley et al. (2022)",
        "title": "Experiences in LP-IoT: EnviSense Deployment of Remotely Reprogrammable Environmental Sensors",
        "doi": "10.1145/3477085.3478988",
        "sensors": "MaxBotix ultrasonic range sensor or SDI-12 pressure transducer; internal T/RH/pressure; battery voltage/current; GPS; LoRa.",
        "materials": "Commercial IP67 enclosure; exact enclosure material not reported.",
        "geometry": "Sealed rectangular logger with four sealed ports and external measurement probes.",
        "setup": "More than 12 nodes at two sites; six months of continuous field measurement discussed.",
        "shows": "Sealed electronics, external probes, remote management, and device-health sensing improve autonomy and diagnosability.",
        "conclusion": "A modular sealed logger can support long deployments when sensing, power, communications, and maintenance are designed together.",
        "importance": "Strong precedent for a two-zone architecture and for treating battery and communication health as measurement-system outputs.",
        "consider": [
            "Separate protected electronics from ambient-air sensors.",
            "Log battery, resets, and communication health alongside environmental data.",
        ],
    },
    {
        "category": "Complete Sensor-Box and Field-System Studies",
        "citation": "Daepp et al. (2022)",
        "title": "Eclipse: An End-to-End Platform for Low-Cost, Hyperlocal Environmental Sensing in Cities",
        "doi": "10.1109/IPSN54338.2022.00010",
        "sensors": "PM2.5; temperature; RH; pressure; optional O3, NO2, SO2, and CO modules.",
        "materials": "Custom weatherproof enclosure; exact material and airflow path not reported.",
        "geometry": "Modular solar-powered, cellular-connected urban sensor box.",
        "setup": "115 devices deployed across Chicago; more than 90% of expected sensor-hours collected during the reported period.",
        "shows": "Large-scale deployment is possible when the platform is modular and expected sensor-hours are treated as an operational metric.",
        "conclusion": "Completeness and remote serviceability are central measures of whether a low-cost box is useful in practice.",
        "importance": "Provides a clear autonomy metric and supports designing reset/service actions that do not require opening the box.",
        "consider": [
            "Report expected versus received sensor-hours.",
            "Evaluate site-specific mounting and solar-panel heating effects.",
        ],
    },
    {
        "category": "Complete Sensor-Box and Field-System Studies",
        "citation": "Lazarescu (2015)",
        "title": "Design and Field Test of a WSN Platform Prototype for Long-Term Environmental Monitoring",
        "doi": "10.3390/s150409481",
        "sensors": "NTC temperature transducer; low-power radio; battery and system-health measurements.",
        "materials": "Protective resin; compact node package; sealed plastic gateway package inside a divided wooden housing.",
        "geometry": "Downward-facing vented node; solar-powered gateway with separated compartments.",
        "setup": "Fifty sensor nodes and two gateways deployed for approximately 2.5 months in an olive field.",
        "shows": "Solar heating, solder fatigue, antenna placement, charger failure, RF attenuation, and faulty-node traffic can limit autonomy.",
        "conclusion": "Theoretical battery life does not establish autonomy; observed mechanical, thermal, RF, and firmware behavior does.",
        "importance": "Directly informs the paper's failure-mode and maintenance-event framework.",
        "consider": [
            "Test antenna clearance and enclosure effects on communication.",
            "Include heartbeats, self-test, battery-state validation, and failure propagation in the deployment plan.",
        ],
    },
    {
        "category": "Radiation-Shield Geometry and Material Studies",
        "citation": "Botero-Valencia et al. (2022)",
        "title": "Design and Implementation of 3-D Printed Radiation Shields for Environmental Sensors",
        "doi": "10.1016/j.ohx.2022.e00267",
        "sensors": "Three identical SHT10 temperature/RH sensors; Vantage Pro reference; Particle Argon logger.",
        "materials": "Orange PLA and white ASA; PVC mount; metal spacers; IP69 electronics box.",
        "geometry": "Six truncated cones: five perforated cones and a solid roof, separated by 20 mm spacers.",
        "setup": "Exposed, PLA-shielded, and ASA-shielded sensors compared outdoors, plus 30/90-day weathering tests.",
        "shows": "Both printed shields reduced radiation error; ASA was favored for outdoor UV resistance and durability.",
        "conclusion": "A modular passive cone-stack shield is a defensible low-cost baseline, but material and color were not fully isolated.",
        "importance": "Provides the most direct literature basis for testing printed passive shields and comparing material/weathering behavior.",
        "consider": [
            "Control color and surface finish when comparing base materials.",
            "Treat cone count, spacing, and open area as independent geometry variables.",
        ],
    },
    {
        "category": "Radiation-Shield Geometry and Material Studies",
        "citation": "Tarara and Hoheisel (2007)",
        "title": "Low-Cost Shielding to Minimize Radiation Errors of Temperature Sensors in the Field",
        "doi": "10.21273/HORTSCI.42.6.1372",
        "sensors": "Same temperature-sensor type used across test and comparison shields.",
        "materials": "Reflective foil insulation with high albedo and low emissivity.",
        "geometry": "Plate, stacked-plate, cone, double-cone, tube, rocket, pagoda, and perforated-tube variants.",
        "setup": "Controlled outdoor comparison under changing solar irradiance and wind.",
        "shows": "Stacked plates performed best; open bottoms admitted reflected radiation; tube shields needed substantial sidewall perforation.",
        "conclusion": "Radiation blocking alone is insufficient: reflected radiation and ventilation area determine passive-shield error.",
        "importance": "Provides the clearest evidence for selecting and parameterizing passive geometry variants.",
        "consider": [
            "Measure vent/open-area ratio and include reflected ground radiation.",
            "Avoid open-bottom and weakly perforated tube designs.",
        ],
    },
    {
        "category": "Radiation-Shield Geometry and Material Studies",
        "citation": "Holden et al. (2013)",
        "title": "Design and Evaluation of an Inexpensive Radiation Shield for Monitoring Surface Air Temperatures",
        "doi": "10.1016/j.agrformet.2013.06.011",
        "sensors": "Low-cost temperature sensors, including LogTag units.",
        "materials": "White corrugated-plastic sheet.",
        "geometry": "Folded Gill-style series of plates.",
        "setup": "Three-site outdoor comparison against commercial passive shields, RAWS, and aspirated ASOS reference.",
        "shows": "A roughly USD 3 shield performed similarly to a commercial passive shield, while low-wind full-sun bias remained.",
        "conclusion": "Simple, inexpensive geometry can be effective enough for distributed sensing, but passive limitations remain.",
        "importance": "Creates a useful low-cost manufacturability benchmark for the lab's enclosure comparisons.",
        "consider": [
            "Compare build repeatability, weathering, and cost alongside accuracy.",
            "Use an aspirated or higher-grade reference to quantify remaining passive bias.",
        ],
    },
    {
        "category": "Radiation-Shield Geometry and Material Studies",
        "citation": "Liu et al. (2023)",
        "title": "Design and Experiments of a Naturally-Ventilated Radiation Shield for Ground Temperature Measurement",
        "doi": "10.3390/atmos14030523",
        "sensors": "Pt100 platinum resistance thermometer.",
        "materials": "Silver-coated aluminum outer plates; black inner surfaces; white-resin internal parts.",
        "geometry": "Multilayer shield with two square plates, annular layers, diversion pipe, eight vents, and truncated-cone reflector.",
        "setup": "CFD-informed design followed by outdoor experimental validation.",
        "shows": "Outer reflectivity, inner surface treatment, venting, and lower reflectors can address direct and reflected radiation together.",
        "conclusion": "Mixed surface finishes and guided natural ventilation offer performance gains, with added fabrication and coating-maintenance complexity.",
        "importance": "Expands the paper beyond polymer choice to include surface finish and reflected-radiation management.",
        "consider": [
            "Treat exterior reflectance and interior finish as separate variables.",
            "Include vent count/area and coating degradation in the decision framework.",
        ],
    },
    {
        "category": "Radiation-Shield Geometry and Material Studies",
        "citation": "García Izquierdo et al. (2024)",
        "title": "COAT Project: Intercomparison of Thermometer Radiation Shields in the Arctic",
        "doi": "10.3390/atmos15070841",
        "sensors": "Forty-one calibrated four-wire Pt100 thermometers; ultrasonic anemometers and reference meteorology.",
        "materials": "Ten commercial shield models; detailed proprietary materials vary.",
        "geometry": "Natural and active shields, including multiplate and Stevenson-screen configurations.",
        "setup": "Fourteen-month Arctic intercomparison with pre/post calibration.",
        "shows": "Shield differences depend strongly on wind and radiation; fan, logger, power, memory, mounting, and storm failures caused data loss.",
        "conclusion": "Long-duration comparison must evaluate operational reliability as well as thermometer error.",
        "importance": "Provides a strong precedent for long-term failure logging and passive-versus-active comparison.",
        "consider": [
            "Log fan operation, storage capacity, power status, and mount condition.",
            "Use redundant probes where possible to distinguish shield effects from sensor failure.",
        ],
    },
    {
        "category": "Radiation-Shield Geometry and Material Studies",
        "citation": "Deford et al. (2025)",
        "title": "Development of a Low-Cost, 3D-Printed, Aspirated Air Temperature Measurement Radiation Shield",
        "doi": "10.1175/JTECH-D-24-0006.1",
        "sensors": "Compact air-temperature sensor; Sunon 5 V fan.",
        "materials": "White PLA; hot glue during testing; fan and exposed circuits.",
        "geometry": "Three-part double-wall 90-degree elbow with insect screen, central supports, and angled fan diffuser.",
        "setup": "Outdoor comparison with passive and commercial aspirated reference shields.",
        "shows": "Low-cost aspiration can achieve strong accuracy, but the fan consumes 680 mW and creates durability and maintenance risks.",
        "conclusion": "Active airflow is a useful accuracy benchmark and a direct test of the accuracy-autonomy tradeoff.",
        "importance": "Provides the clearest candidate active design for comparison with passive lab-box options.",
        "consider": [
            "Measure fan power, flow, and failure behavior.",
            "Replace temporary PLA/glue/exposed circuitry with weatherable materials and permanent assembly for long deployment.",
        ],
    },
    {
        "category": "Radiation-Shield Geometry and Material Studies",
        "citation": "Jin et al. (2026)",
        "title": "Design and Experimental Validation of a High-Accuracy Naturally Ventilated Radiation Shield",
        "doi": "10.3390/atmos17030272",
        "sensors": "Thin-film Pt100 in an 8 mm copper spherical shell.",
        "materials": "CFD comparison of plastic, wood, Fe-Ni alloy, and aluminum; mirror-finish aluminum considered.",
        "geometry": "Two square shading plates with symmetric bowl-cover flow-guiding shrouds and 30 mm spacing.",
        "setup": "CFD material/geometry optimization followed by outdoor validation.",
        "shows": "Material thermal properties and guided passive airflow jointly affect radiation error.",
        "conclusion": "Advanced passive geometry may reduce error without fan power, but complexity and weathering must be assessed.",
        "importance": "Supports evaluating airflow guidance and upward/reflected radiation rather than only plate count.",
        "consider": [
            "Balance modeled accuracy against manufacturability, size, mounting, and coating durability.",
            "Validate CFD-derived designs across multiple weather states.",
        ],
    },
    {
        "category": "Calibration, Drift, and Air-Quality Sensor Studies",
        "citation": "deSouza et al. (2022)",
        "title": "Calibrating Networks of Low-Cost Air Quality Sensors",
        "doi": "10.5194/amt-15-6309-2022",
        "sensors": "Canary-S boxes using Plantower PMS5003 PM sensing with integrated temperature/RH.",
        "materials": "Commercial outdoor box; material and internal airflow not reported.",
        "geometry": "Commercial networked box; detailed inlet/outlet and internal layout not reported.",
        "setup": "Twenty-four school deployments; seven boxes co-located with five FEM reference sites.",
        "shows": "Calibration evaluated at one co-location may fail to transfer across time or space.",
        "conclusion": "Calibration performance must be tested beyond the location and period used to fit the model.",
        "importance": "Directly supports time-separated and, if possible, site-separated validation of the lab box.",
        "consider": [
            "Document PM inlet, fan path, sensor placement, and enclosure material.",
            "Report calibration transfer performance rather than only training/co-location error.",
        ],
    },
    {
        "category": "Calibration, Drift, and Air-Quality Sensor Studies",
        "citation": "Clements et al. (2017)",
        "title": "Low-Cost Air Quality Monitoring Tools: From Research to Practice",
        "doi": "10.3390/s17112478",
        "sensors": "Optical PM, electrochemical gas, metal-oxide gas, and multi-sensor systems discussed generally.",
        "materials": "No single enclosure material reported.",
        "geometry": "No single geometry; air exchange and internal heat are identified as important.",
        "setup": "Workshop summary synthesizing laboratory, co-location, mobile, community, and network experience.",
        "shows": "Co-location, repeated calibration, drift checks, and environmental measurements are common requirements.",
        "conclusion": "Low-cost sensor boxes need a defined quality-assurance workflow rather than a one-time validation.",
        "importance": "Provides broad best-practice justification for calibration and enclosure-condition logging.",
        "consider": [
            "Measure internal heat and air exchange rather than treating the box as neutral.",
            "Place environmental-correction sensors close to pollutant sensors while exposing them to ambient air.",
        ],
    },
    {
        "category": "Calibration, Drift, and Air-Quality Sensor Studies",
        "citation": "Giordano et al. (2021)",
        "title": "From Low-Cost Sensors to High-Quality Data",
        "doi": "10.1016/j.jaerosci.2021.105833",
        "sensors": "Low-cost optical particulate-matter sensors as a class.",
        "materials": "No single enclosure material reported.",
        "geometry": "No single enclosure; packaging and airflow identified as influences.",
        "setup": "Review of laboratory evaluation, co-location, and network deployment studies.",
        "shows": "PM response depends on humidity, aerosol properties, sensor age, site, and packaging/airflow.",
        "conclusion": "PM calibration and enclosure design are deployment-specific and cannot be treated independently.",
        "importance": "Defines environmental covariates and limitations that must be considered if the lab box includes PM sensing.",
        "consider": [
            "Include humidity, inlet airflow, aging, and local aerosol conditions in PM analysis.",
            "Do not infer a universally transferable calibration.",
        ],
    },
    {
        "category": "Calibration, Drift, and Air-Quality Sensor Studies",
        "citation": "Vajs et al. (2021)",
        "title": "Developing Relative Humidity and Temperature Corrections for Low-Cost Sensors Using Machine Learning",
        "doi": "10.3390/s21103338",
        "sensors": "Alphasense gas cells; Bosch BME280; Plantower PMS7003.",
        "materials": "Commercial AQ10x enclosure; material and airflow not reported.",
        "geometry": "Commercial outdoor air-quality station; internal layout not reported.",
        "setup": "One station co-located with a Serbian public monitoring station; LR, MLR, and ML corrections compared.",
        "shows": "Temperature and RH correction can improve low-cost gas/PM measurements, but sensor-specific sensitivity remains.",
        "conclusion": "Interpretable linear and multiple-linear models provide a defensible baseline before more complex methods.",
        "importance": "Directly supports the proposed calibration sequence for the lab box.",
        "consider": [
            "Start with linear and multiple-linear calibration and use a time-separated test period.",
            "Ensure T/RH sensors measure the air reaching the pollutant sensors, not only the electronics-box microclimate.",
        ],
    },
    {
        "category": "Calibration, Drift, and Air-Quality Sensor Studies",
        "citation": "Schneider et al. (2023)",
        "title": "Deployment and Evaluation of a Network of Open Low-Cost Air Quality Sensor Systems",
        "doi": "10.3390/atmos14030540",
        "sensors": "Alphasense NO2, CO, NO, O3; ELT D-300 CO2; OPC-N3 PM; T/RH/pressure.",
        "materials": "AirSensEUR boxes; enclosure material not reported.",
        "geometry": "Complete boxes co-located together, then distributed to roadside sites.",
        "setup": "Five-week pre-deployment co-location, roadside deployment, and post-deployment evaluation.",
        "shows": "High RH affected optical PM; unit-to-unit differences, drift, and failures required pre/post deployment evaluation.",
        "conclusion": "Pre- and post-deployment co-location improves detection of drift and field-induced change.",
        "importance": "Supports comparing multiple lab boxes and checking calibration after deployment.",
        "consider": [
            "Perform pre/post co-location if deployment duration permits.",
            "Treat humidity, unit variability, and multi-sensor maintenance as primary factors.",
        ],
    },
    {
        "category": "Calibration, Drift, and Air-Quality Sensor Studies",
        "citation": "Diez et al. (2024)",
        "title": "Long-Term Evaluation of Commercial Air Quality Sensors: The QUANT Study",
        "doi": "10.5194/amt-17-3809-2024",
        "sensors": "Forty-three commercial devices containing 119 gas and 118 PM sensors.",
        "materials": "Commercial protective cases; materials and internal layouts generally proprietary.",
        "geometry": "Multiple commercial black-box systems.",
        "setup": "Three-year evaluation across UK urban environments, including relocation and uncertainty analysis.",
        "shows": "Performance varied substantially between systems; relocation and manufacturer correction changes affected results.",
        "conclusion": "Long-term validation requires documentation of hardware, firmware, processing, location, and uncertainty.",
        "importance": "Supports treating the lab box as a transparent system and evaluating more than RMSE alone.",
        "consider": [
            "Document every hardware, firmware, and calibration change.",
            "Evaluate inter-device precision, uncertainty, relocation, and long-term stability.",
        ],
    },
    {
        "category": "Calibration, Drift, and Air-Quality Sensor Studies",
        "citation": "Winter et al. (2025)",
        "title": "Sustained Performance of Low-Cost Air Quality Sensors in Long-Term Deployments",
        "doi": "10.1021/acssensors.5c00566",
        "sensors": "Alphasense NO2-B43F, Ox-B431, NO-B4, and CO-B4 electrochemical sensors.",
        "materials": "BEACO2N node enclosure not described in detail.",
        "geometry": "Dense-network node packaging; geometry not the focus.",
        "setup": "Three-year San Francisco Bay Area network evaluation.",
        "shows": "Calibrated electrochemical performance can remain stable for years, while mechanical/electrical failures may occur first.",
        "conclusion": "Recalibration interval and physical-maintenance interval are separate engineering quantities.",
        "importance": "Strengthens the paper's focus on identifying the true limiting factor rather than assuming sensor drift dominates.",
        "consider": [
            "Track recalibration need separately from physical interventions.",
            "Record board, connector, power, and enclosure failures alongside sensing-element behavior.",
        ],
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=10.5, bold=False, italic=False, color=BLACK, font="Calibri") -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_para(doc, text="", *, style=None, bold_lead=None, color=BLACK, size=None, after=6, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.keep_with_next = keep
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run(r1, size=size or 10.5, bold=True, color=color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run(r2, size=size or 10.5, color=color)
    else:
        r = p.add_run(text)
        set_run(r, size=size or 10.5, color=color)
    return p


def add_bullet(doc, text, *, color=BLACK, after=3):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_run(r, size=10.2, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_key_facts_table(doc, profile):
    rows = [
        ("Sensors", profile["sensors"]),
        ("Materials", profile["materials"]),
        ("Geometry", profile["geometry"]),
        ("Experimental setup", profile["setup"]),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(1.35)
        cells[1].width = Inches(5.05)
        set_cell_shading(cells[0], LIGHT_BLUE)
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        set_run(p0.add_run(label), size=9.3, bold=True, color=DARK)
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run(p1.add_run(value), size=9.3, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_profile(doc, profile, number):
    add_heading(doc, f"{number}. {profile['citation']}", 2)
    p = add_para(doc, profile["title"], color=DARK, size=11, after=3, keep=True)
    p.runs[0].italic = True
    add_para(doc, f"DOI: https://doi.org/{profile['doi']}", color=MID_GRAY, size=8.8, after=6, keep=True)
    add_key_facts_table(doc, profile)
    add_para(doc, f"What it shows: {profile['shows']}", bold_lead="What it shows: ", after=4)
    add_para(doc, f"Brief conclusion: {profile['conclusion']}", bold_lead="Brief conclusion: ", after=4)
    add_para(doc, f"Importance to this research: {profile['importance']}", bold_lead="Importance to this research: ", after=3)
    if profile["consider"]:
        add_para(doc, "Things to consider moving forward", bold_lead="Things to consider moving forward", after=2, keep=True)
        for item in profile["consider"]:
            add_bullet(doc, item, after=2)


def add_summary_table(doc):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    widths = [0.45, 2.15, 1.75, 2.05]
    headers = ["#", "Paper", "Primary contribution", "Most relevant design issue"]
    for index, (cell, width, header) in enumerate(zip(table.rows[0].cells, widths, headers)):
        cell.width = Inches(width)
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell, top=100, bottom=100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(header), size=8.6, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    contribution_map = {
        "Complete Sensor-Box and Field-System Studies": "Complete system / autonomy",
        "Radiation-Shield Geometry and Material Studies": "Shield geometry / material",
        "Calibration, Drift, and Air-Quality Sensor Studies": "Calibration / drift",
    }
    for i, profile in enumerate(profiles, 1):
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = Inches(width)
            set_cell_margins(cell, top=70, bottom=70)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        values = [
            str(i),
            profile["citation"],
            contribution_map[profile["category"]],
            profile["conclusion"],
        ]
        for col, (cell, value) in enumerate(zip(cells, values)):
            if i % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            set_run(p.add_run(value), size=7.8, color=BLACK)


def build_document():
    OUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 12.5, BLUE, 10, 5),
        ("Heading 3", 11, DARK, 7, 3),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Number"]:
        styles[name].font.name = "Calibri"
        styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        styles[name].font.size = Pt(10.2)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    set_run(hp.add_run("Outdoor Multi-Sensor Box | Literature Synthesis"), size=8.5, color=MID_GRAY)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    set_run(fp.add_run("Research working report  |  Page "), size=8.5, color=MID_GRAY)
    add_page_number(fp)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(78)
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("LITERATURE SYNTHESIS"), size=12, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    set_run(p.add_run("Accuracy, Calibration, Autonomy, and Physical Integration of Low-Cost Outdoor Sensor Boxes"), size=24, bold=True, color=DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    set_run(p.add_run("Summary of the 19 research sources currently included in the Enclosure-Research repository"), size=12, italic=True, color=MID_GRAY)

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row, (label, value) in zip(
        table.rows,
        [
            ("Prepared for", "Principal Investigator review"),
            ("Working paper", "Calibration, Ambient Conditions, and Integration Effects on the Accuracy and Autonomy of a Low-Cost Outdoor Multi-Sensor Box"),
            ("Date", "June 11, 2026"),
        ],
    ):
        row.cells[0].width = Inches(1.25)
        row.cells[1].width = Inches(4.95)
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            set_cell_margins(cell, top=110, bottom=110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_run(row.cells[0].paragraphs[0].add_run(label), size=9.5, bold=True, color=DARK)
        set_run(row.cells[1].paragraphs[0].add_run(value), size=9.5, color=BLACK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("Purpose"), size=11, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.55)
    p.paragraph_format.right_indent = Inches(0.55)
    p.paragraph_format.space_after = Pt(0)
    set_run(
        p.add_run(
            "This report connects each literature source to the proposed study, summarizes the reported sensors, materials, and geometries, and identifies evidence-supported considerations for the next experimental phase."
        ),
        size=10.5,
        color=DARK,
    )

    doc.add_page_break()
    add_heading(doc, "Executive Summary", 1)
    add_para(
        doc,
        "The literature supports framing the lab's sensor box as an integrated measurement instrument rather than as a collection of individual sensors or an enclosure-only design. Field accuracy depends on calibration, but it also depends on airflow, direct and reflected radiation, sensor placement, electronics self-heating, moisture protection, power, communication, storage, and maintenance.",
        after=7,
    )
    add_para(
        doc,
        "The strongest recurring result is that passive shields fail most noticeably during high-solar, low-wind conditions. Stacked plates, overlapping louvers, reflective surfaces, and separated electronics improve performance without consuming power. Active aspiration can further reduce airflow-dependent bias, but the fan directly reduces autonomy and adds a failure mode. Long-duration deployments consistently show that connectors, mounts, batteries, charging, storage, communications, and firmware may fail before the sensing element itself.",
        after=7,
    )
    add_heading(doc, "Implications for the Proposed Paper", 2)
    for item in [
        "Begin with the current lab box as an unchanged baseline and compare it with reference instruments.",
        "Separate the sealed electronics/power region from sensors that require ambient airflow.",
        "Evaluate raw and calibrated accuracy, but do not use calibration to conceal avoidable enclosure heating or airflow problems.",
        "Report autonomy using runtime before intervention, expected versus received samples, uptime, maintenance events, and failure types.",
        "Use a small, justified design comparison: current box, strong passive shield, and optional aspirated benchmark.",
        "Treat material, color/finish, geometry, airflow, sealing, and sensor-to-electronics distance as measurable factors.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Cross-Paper Findings", 2)
    add_para(doc, "Sensor selection", bold_lead="Sensor selection", after=2, keep=True)
    add_bullet(doc, "Digital temperature/RH sensors are compact and easy to integrate, but their readings are strongly affected by shielding, airflow, condensation, and placement.")
    add_bullet(doc, "Optical PM and electrochemical gas sensors require environmental correction, repeated calibration, and controlled exposure to ambient air.")
    add_bullet(doc, "Reference-grade Pt100 studies provide strong shield-comparison methods, even when the final lab sensor is lower cost.")
    add_para(doc, "Material and geometry", bold_lead="Material and geometry", after=2, keep=True)
    add_bullet(doc, "White/off-white ASA is the best-supported printed outdoor material; PLA is useful for prototypes but requires durability caution.")
    add_bullet(doc, "Reflective foil and white corrugated plastic provide very low-cost geometry benchmarks.")
    add_bullet(doc, "Passive stacked plates generally outperform open or weakly ventilated tube geometries.")
    add_bullet(doc, "IP-rated boxes are appropriate for electronics, but ambient sensors should be external or in a separately ventilated region.")
    add_para(doc, "Autonomy and reliability", bold_lead="Autonomy and reliability", after=2, keep=True)
    add_bullet(doc, "Battery life alone is not autonomy; physical intervention can be triggered by water, corrosion, memory, fan, firmware, RF, connector, or mounting failures.")
    add_bullet(doc, "Pre/post co-location and device-health logging improve the ability to identify drift and limiting failures.")

    doc.add_page_break()
    add_heading(doc, "Literature at a Glance", 1)
    add_para(doc, "The table below identifies the primary contribution of each source. Detailed source profiles follow.", color=MID_GRAY, size=9.5, after=8)
    add_summary_table(doc)

    current_category = None
    number = 0
    for profile in profiles:
        if profile["category"] != current_category:
            doc.add_page_break()
            current_category = profile["category"]
            add_heading(doc, current_category, 1)
            if current_category == "Complete Sensor-Box and Field-System Studies":
                add_para(doc, "These papers show how sensors, packaging, power, communication, and maintenance behave as one deployed system.", color=MID_GRAY, size=9.5)
            elif current_category == "Radiation-Shield Geometry and Material Studies":
                add_para(doc, "These papers provide the strongest evidence for selecting and comparing passive and active shield designs.", color=MID_GRAY, size=9.5)
            else:
                add_para(doc, "These papers define calibration, drift, environmental correction, and long-term validation requirements.", color=MID_GRAY, size=9.5)
        number += 1
        add_profile(doc, profile, number)

    doc.add_page_break()
    add_heading(doc, "Recommended Experimental Framing", 1)
    add_para(
        doc,
        "The literature supports an engineering comparison centered on the lab's existing box. The primary result should be identification of the limiting factor: sensor calibration, physical integration, power, firmware/data handling, or maintenance.",
    )
    add_heading(doc, "Recommended Initial Designs", 2)
    designs = [
        ("Current lab box", "Unmodified baseline that establishes present accuracy, autonomy, and failure modes."),
        ("Strong passive option", "White reflective stacked-plate or truncated-cone shield with electronics isolated in an IP-rated compartment."),
        ("Active benchmark", "Aspirated shield used only if the power budget permits; records fan power and fan-operating status."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell, width, value in zip(table.rows[0].cells, [1.75, 4.65], ["Design", "Purpose"]):
        cell.width = Inches(width)
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell, top=100, bottom=100)
        set_run(cell.paragraphs[0].add_run(value), size=9, bold=True, color=WHITE)
    for name, purpose in designs:
        cells = table.add_row().cells
        for cell, width in zip(cells, [1.75, 4.65]):
            cell.width = Inches(width)
            set_cell_margins(cell)
        set_run(cells[0].paragraphs[0].add_run(name), size=9.3, bold=True, color=DARK)
        set_run(cells[1].paragraphs[0].add_run(purpose), size=9.3)

    add_heading(doc, "Measurements Needed", 2)
    for item in [
        "Reference-matched sensor readings and timestamps.",
        "Solar exposure or radiation proxy, wind, rain, ambient temperature, and RH.",
        "Battery voltage/current, fan state if used, resets, storage and communication health.",
        "Expected and received sample counts by sensor channel.",
        "Maintenance, calibration, physical inspection, and failure-event log.",
        "Material, color/finish, geometry dimensions, vent/open-area ratio, and sensor placement.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Questions to Resolve Before Deployment", 2)
    for item in [
        "What variables must the lab box measure, and what reference instruments are available for each?",
        "What calibrated-error, completeness, and autonomy thresholds define a useful deployment?",
        "Which current component is most likely to limit performance: sensor, enclosure, power, firmware, or maintenance?",
        "Can the lab fabricate ASA reliably, or should PETG/coated PETG and purchased IP-rated boxes be the practical candidates?",
        "Is active aspiration justified by the expected accuracy improvement and available power budget?",
        "How long must the deployment run to capture representative solar, wind, humidity, rain, and temperature conditions?",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Overall Conclusion", 1)
    add_para(
        doc,
        "Across the 19 sources, the strongest evidence supports a two-zone outdoor sensor-box architecture: protected electronics and power in a sealed, serviceable enclosure, with ambient sensors placed in an external or separately ventilated shield. A light-colored passive stacked-plate or cone shield is the most defensible low-power baseline. Active aspiration is valuable as an accuracy benchmark but must be evaluated against its power and maintenance cost. The paper should therefore report calibration improvement and physical-design performance together, while using observed intervention-free runtime and data completeness to define autonomy.",
        after=8,
    )
    add_para(
        doc,
        "The final contribution should be a practical decision framework showing which design factors most affect the lab's sensor box and which changes provide measurable improvement without unnecessary complexity.",
        after=8,
    )

    add_heading(doc, "Reference List", 1)
    for profile in profiles:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.05
        set_run(p.add_run(f"{profile['citation']}. "), size=9, bold=True, color=DARK)
        set_run(p.add_run(f"{profile['title']}. "), size=9, italic=True, color=BLACK)
        set_run(p.add_run(f"https://doi.org/{profile['doi']}"), size=9, color=BLUE)

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_document()
